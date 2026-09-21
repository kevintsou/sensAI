# -*- coding: utf-8 -*-
"""sensAI 主管簡報產生器。"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

SHOTS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets")
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "sensAI_報告.pptx")

NAVY   = RGBColor(0x10, 0x2A, 0x43)
NAVY2  = RGBColor(0x1C, 0x44, 0x6B)
TEAL   = RGBColor(0x11, 0x9D, 0xA4)
AMBER  = RGBColor(0xD1, 0x86, 0x16)
GREY   = RGBColor(0x5B, 0x6B, 0x7B)
LIGHT  = RGBColor(0xF2, 0xF5, 0xF8)
LINE   = RGBColor(0xD8, 0xDF, 0xE6)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
DARK   = RGBColor(0x1B, 0x26, 0x33)

FONT = "微軟正黑體"
MONO = "Consolas"

W, H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


def tb(slide, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    s = slide.shapes.add_textbox(x, y, w, h)
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return tf


def para(tf, text, size=16, color=DARK, bold=False, space_after=6,
         first=False, align=None, font=FONT, space_before=0, line=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    if align is not None:
        p.alignment = align
    if line:
        p.line_spacing = line
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font
    return p


def rect(slide, x, y, w, h, fill=None, line_col=None, line_w=1.0, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line_col is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line_col; s.line.width = Pt(line_w)
    s.shadow.inherit = False
    return s


def header(slide, kicker, title):
    """每頁標準頁首：左側色條 + 小標 + 主標 + 分隔線。"""
    rect(slide, Inches(0), Inches(0), Inches(0.13), H, fill=NAVY)
    tf = tb(slide, Inches(0.62), Inches(0.42), Inches(11.9), Inches(0.3))
    para(tf, kicker, size=12, color=TEAL, bold=True, first=True, space_after=0)
    tf = tb(slide, Inches(0.62), Inches(0.72), Inches(11.9), Inches(0.55))
    para(tf, title, size=28, color=NAVY, bold=True, first=True, space_after=0)
    rect(slide, Inches(0.62), Inches(1.42), Inches(12.1), Emu(11430), fill=LINE)


def footer(slide, n):
    tf = tb(slide, Inches(0.62), Inches(6.95), Inches(6.0), Inches(0.28))
    para(tf, "sensAI ｜ 韌體 AI Code Review ｜ v0.5.1", size=10, color=GREY, first=True, space_after=0)
    tf = tb(slide, Inches(11.6), Inches(6.95), Inches(1.1), Inches(0.28), align=PP_ALIGN.RIGHT)
    para(tf, str(n), size=10, color=GREY, first=True, space_after=0, align=PP_ALIGN.RIGHT)


def new(kicker, title, n):
    s = prs.slides.add_slide(BLANK)
    header(s, kicker, title)
    footer(s, n)
    return s


def stat(slide, x, y, w, value, label, note="", vcolor=NAVY, h=Inches(1.32)):
    """量化數字方塊。"""
    rect(slide, x, y, w, h, fill=LIGHT)
    rect(slide, x, y, Emu(34290), h, fill=TEAL)
    tf = tb(slide, x + Inches(0.22), y + Inches(0.14), w - Inches(0.34), Inches(0.52))
    para(tf, value, size=30, color=vcolor, bold=True, first=True, space_after=0)
    tf = tb(slide, x + Inches(0.22), y + Inches(0.68), w - Inches(0.34), Inches(0.55))
    para(tf, label, size=13, color=DARK, bold=True, first=True, space_after=2)
    if note:
        para(tf, note, size=10.5, color=GREY, space_after=0)


def bullets(slide, x, y, w, items, size=15, gap=9, bullet_col=TEAL):
    tf = tb(slide, x, y, w, Inches(0.4))
    for i, it in enumerate(items):
        if isinstance(it, tuple):
            head, body = it
            p = para(tf, head, size=size, color=NAVY, bold=True,
                     first=(i == 0), space_after=2, space_before=(0 if i == 0 else gap))
            para(tf, body, size=size - 1.5, color=GREY, space_after=0, line=1.25)
        else:
            para(tf, "・" + it, size=size, color=DARK, first=(i == 0),
                 space_after=gap, line=1.25)
    return tf


def pic(slide, path, x, y, w=None, h=None, border=True):
    if w is not None:
        p = slide.shapes.add_picture(path, x, y, width=w)
    else:
        p = slide.shapes.add_picture(path, x, y, height=h)
    if border:
        b = rect(slide, p.left - Emu(9525), p.top - Emu(9525),
                 p.width + Emu(19050), p.height + Emu(19050),
                 fill=None, line_col=RGBColor(0x8A, 0x97, 0xA5), line_w=0.75)
        # 邊框放到圖片後面
        slide.shapes._spTree.remove(b._element)
        slide.shapes._spTree.insert(list(slide.shapes._spTree).index(p._element), b._element)
    return p


def caption(slide, x, y, w, text):
    tf = tb(slide, x, y, w, Inches(0.3))
    para(tf, text, size=10.5, color=GREY, first=True, space_after=0)


# ───────────────────────────────── 1 封面
s = prs.slides.add_slide(BLANK)
rect(s, Inches(0), Inches(0), W, H, fill=NAVY)
rect(s, Inches(0), Inches(0), Inches(0.18), H, fill=TEAL)
tf = tb(s, Inches(1.1), Inches(2.18), Inches(11), Inches(0.4))
para(tf, "韌體開發效率提案", size=15, color=TEAL, bold=True, first=True, space_after=0)
tf = tb(s, Inches(1.1), Inches(2.62), Inches(11), Inches(1.2))
para(tf, "sensAI", size=60, color=WHITE, bold=True, first=True, space_after=0)
tf = tb(s, Inches(1.1), Inches(3.78), Inches(11), Inches(0.9))
para(tf, "ARM / Andes AndeStar V5 韌體專用的 AI Code Review",
     size=22, color=WHITE, first=True, space_after=4)
para(tf, "把團隊的硬體知識寫成規則，在存檔當下就攔下缺陷", size=15, color=RGBColor(0xA8, 0xC4, 0xD8))
rect(s, Inches(1.1), Inches(5.0), Inches(2.4), Emu(22860), fill=TEAL)
tf = tb(s, Inches(1.1), Inches(5.32), Inches(11), Inches(0.9))
para(tf, "VS Code 擴充　v0.5.1　·　已在實機專案驗證", size=13,
     color=RGBColor(0xA8, 0xC4, 0xD8), first=True, space_after=4)
para(tf, "報告人：kevin.tsou　·　2026 年 9 月", size=13, color=RGBColor(0xA8, 0xC4, 0xD8))

# ───────────────────────────────── 2 一頁摘要
s = new("EXECUTIVE SUMMARY", "一頁摘要", 2)
stat(s, Inches(0.62), Inches(1.72), Inches(2.85), "7 個月",
     "實戰命中的潛伏缺陷", "一行被註解掉的 return，存檔後 66 秒攔下", vcolor=AMBER)
stat(s, Inches(3.67), Inches(1.72), Inches(2.85), "66 秒",
     "單檔完整審查耗時", "自動觸發、背景執行，不影響編輯")
stat(s, Inches(6.72), Inches(1.72), Inches(2.85), "10 條",
     "團隊硬體規則", "DMA cache／W1C／ISR／ABI，隨專案進版控")
stat(s, Inches(9.77), Inches(1.72), Inches(2.95), "0 元",
     "外購授權成本", "自行開發，96 項自動化測試全數通過")

rect(s, Inches(0.62), Inches(3.34), Inches(12.1), Inches(1.02), fill=LIGHT)
tf = tb(s, Inches(0.92), Inches(3.5), Inches(11.5), Inches(0.75))
para(tf, "問題", size=13, color=TEAL, bold=True, first=True, space_after=3)
para(tf, "韌體缺陷的代價不在「寫錯」，在「三天後才在板子上發現」。DMA cache、W1C 暫存器、ISR 安全性、"
         "組語 ABI 這類問題屬於各專案的硬體知識，通用 AI 工具與 code review 都不一定看得出來。",
     size=13.5, color=DARK, space_after=0, line=1.25)

rect(s, Inches(0.62), Inches(4.5), Inches(5.9), Inches(2.18), fill=None, line_col=LINE)
tf = tb(s, Inches(0.92), Inches(4.68), Inches(5.4), Inches(1.9))
para(tf, "做法", size=13, color=TEAL, bold=True, first=True, space_after=6)
for t in ["把團隊的硬體慣例寫成自然語言規則，隨專案進版控",
          "存檔且相對 git HEAD 有改動時自動審查，無需人工操作",
          "每則意見必須交代觸發條件、後果、程式碼依據",
          "只提供意見、不畫紅線、不自動改碼 —— 判斷權留在工程師手上"]:
    para(tf, "・" + t, size=12.5, color=DARK, space_after=5, line=1.2)

rect(s, Inches(6.82), Inches(4.5), Inches(5.9), Inches(2.18), fill=None, line_col=LINE)
tf = tb(s, Inches(7.12), Inches(4.68), Inches(5.4), Inches(1.9))
para(tf, "目前狀態與請求", size=13, color=TEAL, bold=True, first=True, space_after=6)
for t in ["v0.5.1 已可安裝使用，實機專案已驗證出真實缺陷",
          "開發至今 14 天、45 次提交、3,603 行程式碼",
          "請求：指定 1〜2 個試行專案，由該專案負責人補齊規則",
          "請求：核配一組共用模型端點（CCR）供團隊連線"]:
    para(tf, "・" + t, size=12.5, color=DARK, space_after=5, line=1.2)

# ───────────────────────────────── 3 問題
s = new("PROBLEM", "為什麼韌體特別需要這件事", 3)
bullets(s, Inches(0.62), Inches(1.78), Inches(6.0), [
    ("缺陷的發現點離寫錯的地方很遠",
     "DMA 搬到過期資料、W1C 一次清掉別人的旗標，這些問題在編譯階段完全合法，"
     "通常要等到板子上出現偶發異常才被察覺，重現成本極高。"),
    ("通用 AI 工具不知道你的 SoC",
     "「D-cache 是 write-back」「STATUS 是 write-one-to-clear」「_isr 結尾的變數不算共享」—— "
     "這些是專案知識，不在任何模型的訓練資料裡。"),
    ("人工 code review 會漏",
     "右側是實際案例：一行被註解掉的 return，通過了 code review，在主線存活 7 個月。"),
])

rect(s, Inches(7.0), Inches(1.78), Inches(5.72), Inches(4.5), fill=LIGHT)
tf = tb(s, Inches(7.3), Inches(2.0), Inches(5.12), Inches(0.4))
para(tf, "實際案例：phal_sys.c", size=15, color=NAVY, bold=True, first=True, space_after=10)
cb = rect(s, Inches(7.3), Inches(2.5), Inches(5.12), Inches(1.38), fill=WHITE, line_col=LINE)
tf = tb(s, Inches(7.5), Inches(2.66), Inches(4.8), Inches(1.1))
para(tf, "U32 phal_sys_get_all_gpio_input_value(void)", size=11.5, color=DARK,
     first=True, space_after=2, font=MONO)
para(tf, "{    ...", size=11.5, color=GREY, space_after=2, font=MONO)
para(tf, "    //return ulInput;", size=11.5, color=AMBER, bold=True, space_after=2, font=MONO)
para(tf, "}", size=11.5, color=GREY, space_after=0, font=MONO)
tf = tb(s, Inches(7.3), Inches(4.05), Inches(5.12), Inches(2.1))
for t in ["宣告回傳 U32，卻沒有任何 return",
          "整個 GPIO 讀取迴圈會被編譯器整段優化掉",
          "呼叫端拿到的是暫存器殘值 —— 所有依賴 GPIO 輸入的判斷都建立在錯誤資料上",
          "git blame：7 個月前提交，至今仍在主線"]:
    para(tf, "・" + t, size=12.5, color=DARK, first=(t.startswith("宣告")),
         space_after=6, line=1.2)

# ───────────────────────────────── 4 運作方式
s = new("HOW IT WORKS", "運作方式：六個步驟，全自動", 4)
steps = [
    ("1", "存檔", "偵測 .c / .h / .s / .S\n相對 git HEAD 無改動即跳過"),
    ("2", "組上下文", "解析專案內 #include\n深度 2、上限 120 KB"),
    ("3", "套規則", "依語言挑選 rules.yaml\n組語另外注入 ABI 事實"),
    ("4", "模型審查", "經 Claude Code Router\n以 tool use 取回結構化結果"),
    ("5", "過濾", "依據引用不到實際識別字者\n直接濾除，不顯示"),
    ("6", "面板呈現", "側欄列出意見\n不畫錯誤紅線"),
]
bx, by, bw, bh = Inches(0.62), Inches(1.86), Inches(1.88), Inches(2.0)
gapx = Inches(2.02)
for i, (num, title, body) in enumerate(steps):
    x = bx + gapx * i
    rect(s, x, by, bw, bh, fill=WHITE, line_col=LINE)
    rect(s, x, by, bw, Inches(0.07), fill=TEAL)
    tf = tb(s, x + Inches(0.18), by + Inches(0.25), bw - Inches(0.36), Inches(0.35))
    para(tf, num, size=20, color=TEAL, bold=True, first=True, space_after=0)
    tf = tb(s, x + Inches(0.18), by + Inches(0.68), bw - Inches(0.36), Inches(1.1))
    para(tf, title, size=15, color=NAVY, bold=True, first=True, space_after=5)
    para(tf, body, size=10.5, color=GREY, space_after=0, line=1.22)
    if i < 5:
        tfa = tb(s, x + bw + Emu(9525), by + Inches(0.78), Inches(0.14), Inches(0.3))
        para(tfa, "›", size=18, color=RGBColor(0xB0, 0xBC, 0xC8), first=True, space_after=0)

rect(s, Inches(0.62), Inches(4.2), Inches(12.1), Inches(2.1), fill=LIGHT)
tf = tb(s, Inches(0.95), Inches(4.42), Inches(11.5), Inches(0.35))
para(tf, "三層節流：不會因為頻繁存檔就狂送請求", size=15, color=NAVY, bold=True,
     first=True, space_after=10)
cols = [("去抖動 1,000 ms", "存檔後等 1 秒沒有新存檔才送出。打字期間幾乎不送任何請求。"),
        ("合併重跑", "同一檔案同時只跑一輪；期間進來的存檔合併成一次補跑，用最新內容。"),
        ("連續觸發時降級", "連打期間只看改動的行，停下來後自動補做完整審查 —— 是延後，不是放棄。")]
for i, (t, b) in enumerate(cols):
    x = Inches(0.95) + Inches(3.92) * i
    tfx = tb(s, x, Inches(4.9), Inches(3.55), Inches(1.2))
    para(tfx, t, size=13, color=TEAL, bold=True, first=True, space_after=4)
    para(tfx, b, size=12, color=DARK, space_after=0, line=1.25)

# ───────────────────────────────── 5 規則即資產
s = new("KEY DESIGN 1", "差異化不在 AI，在規則", 5)
pic(s, os.path.join(SHOTS, "rules.png"), Inches(0.62), Inches(1.78), w=Inches(7.5))
caption(s, Inches(0.62), Inches(5.95), Inches(7.5),
        "實機畫面：.sensai/rules.yaml —— 規則以自然語言撰寫，隨專案進版控，git pull 即生效")
bullets(s, Inches(8.45), Inches(1.82), Inches(4.3), [
    ("規則是團隊資產，不是工具的一部分",
     "規則放在專案 repo，不隨擴充散布，也不會外流。人員異動時，硬體慣例留在版控裡。"),
    ("不需要 DSL、不需要 AST matcher",
     "判斷者是模型，規則直接用中文寫。一組正反範例對準確度的提升，遠大於把規則寫更長。"),
    ("except 讓「不該報什麼」同案管理",
     "誤報的修正方式是改規則，不是關掉工具 —— 規則會越用越準。"),
], size=14, gap=12)
rect(s, Inches(8.45), Inches(5.48), Inches(4.27), Inches(0.98), fill=LIGHT)
tf = tb(s, Inches(8.7), Inches(5.66), Inches(3.8), Inches(0.7))
para(tf, "目前 10 條規則", size=13, color=NAVY, bold=True, first=True, space_after=3)
para(tf, "error 8 條 · warning 2 條 ｜ 另內建 2 組架構 ABI 事實（riscv32-andes-v5、armv7e-m）",
     size=11, color=GREY, space_after=0, line=1.2)

# ───────────────────────────────── 6 防幻覺
s = new("KEY DESIGN 2", "每則意見都要說得出「為什麼」", 6)
tf = tb(s, Inches(0.62), Inches(1.78), Inches(12.1), Inches(0.5))
para(tf, "AI 最大的風險是講得頭頭是道但其實是錯的。sensAI 用結構強制模型交代清楚 —— "
         "說不出具體失效情境的意見，模型自己就不會提。",
     size=14.5, color=DARK, first=True, space_after=0, line=1.25)

fields = [("觸發條件", "什麼情況下會出事", "「編譯器的控制流程走到函式結尾時…」", TEAL),
          ("後果", "會造成什麼影響", "「呼叫端拿到暫存器殘值，判斷建立在錯誤資料上」", AMBER),
          ("依據", "引用檔案中實際存在的行號與識別字", "「第 315 行 //return ulInput；第 306 行函式簽名」", NAVY2)]
y = Inches(2.62)
for t, d, ex, c in fields:
    rect(s, Inches(0.62), y, Inches(7.2), Inches(1.1), fill=WHITE, line_col=LINE)
    rect(s, Inches(0.62), y, Emu(57150), Inches(1.1), fill=c)
    tfx = tb(s, Inches(1.0), y + Inches(0.16), Inches(6.6), Inches(0.85))
    para(tfx, t + "　—　" + d, size=14, color=NAVY, bold=True, first=True, space_after=4)
    para(tfx, ex, size=12, color=GREY, space_after=0, line=1.2)
    y = y + Inches(1.25)

rect(s, Inches(8.1), Inches(2.62), Inches(4.62), Inches(3.6), fill=LIGHT)
tf = tb(s, Inches(8.42), Inches(2.85), Inches(4.0), Inches(3.2))
para(tf, "三道防線", size=15, color=NAVY, bold=True, first=True, space_after=12)
for i, (t, b) in enumerate([
        ("① 結構約束", "欄位填不滿就提不出意見，「建議加強錯誤處理」這類正確但無用的話會自己消失。"),
        ("② 事後過濾", "依據引用不到檔案中真實存在的識別字者，直接濾除、不顯示。"),
        ("③ 呈現方式", "結果只進側欄，不畫編輯器紅線 —— 紅線是斷言，AI 給不起那個確定性。")]):
    para(tf, t, size=13, color=TEAL, bold=True, space_after=3, space_before=(0 if i == 0 else 10))
    para(tf, b, size=11.5, color=DARK, space_after=0, line=1.25)

# ───────────────────────────────── 7 實機畫面
s = new("DEMO", "實機畫面：存檔即審查", 7)
pic(s, os.path.join(SHOTS, "panel.png"), Inches(0.62), Inches(1.78), w=Inches(8.3))
caption(s, Inches(0.62), Inches(6.82), Inches(8.3),
        "版面示意：以 src/panel.ts 的實際樣式重繪，意見內容取自 examples/uart_dma.c 檔尾列出的預期結果")
bullets(s, Inches(9.5), Inches(1.78), Inches(3.25), [
    ("左：編輯器", "行號、程式碼原封不動 —— 不覆寫、不自動修改。"),
    ("右：sensAI 側欄", "檔名、意見數、附帶 header 數、耗時、濾除數、產出時間，一行交代完。"),
    ("每則意見", "嚴重度、行號（可點擊跳轉）、命中的規則 id，再加觸發條件／後果／依據。"),
    ("兩個按鈕", "「釘選」把重要意見固定在頂部、跨重啟保留；「這是誤報」本機立即靜音並留下記錄供調規則。"),
], size=13, gap=11)

# ───────────────────────────────── 8 實戰案例
s = new("RESULT", "實戰驗證：真實專案的真實缺陷", 8)
pic(s, os.path.join(SHOTS, "real.png"), Inches(0.62), Inches(1.74), w=Inches(8.3))
caption(s, Inches(0.62), Inches(5.35), Inches(8.3),
        "實機畫面：ws_fw / ps5032 專案 phal_sys.c —— 存檔後自動審查，一則 error 命中第 315 行")
stat(s, Inches(9.2), Inches(1.74), Inches(3.52), "7 個月", "缺陷潛伏時間",
     "git blame：7 months ago", vcolor=AMBER, h=Inches(1.18))
stat(s, Inches(9.2), Inches(3.04), Inches(3.52), "65.9 秒", "從存檔到看見結果",
     "附帶 18 個專案 header", h=Inches(1.18))
stat(s, Inches(9.2), Inches(4.34), Inches(3.52), "1 則", "意見數（非洗版）",
     "同次審查未產生任何誤報", h=Inches(1.18))
rect(s, Inches(0.62), Inches(5.72), Inches(8.3), Inches(0.95), fill=LIGHT)
tf = tb(s, Inches(0.92), Inches(5.88), Inches(7.7), Inches(0.7))
para(tf, "這個缺陷的意義", size=12.5, color=TEAL, bold=True, first=True, space_after=3)
para(tf, "它不是艱深的時序問題，而是一行被註解掉的 return —— 編譯器只給警告、code review 沒攔下、"
         "在主線活了 7 個月。這正是「人會漏、工具該補」的位置。",
     size=12.5, color=DARK, space_after=0, line=1.25)

# ───────────────────────────────── 9 資安
s = new("SECURITY", "資安與治理：原始碼不會亂跑", 9)
items = [
    ("排除清單（never_send）", "以 glob 指定機密路徑。受審檔案或它引用到的任何 header 命中，"
     "整次審查直接跳過，不外送任何內容。", "src/secure/**　**/crypto/**"),
    ("稽核日誌（audit_log）", "每次對外請求都寫一筆記錄，含結果。v0.5.1 起改為「送出即記錄」—— "
     "失敗的請求也留痕，因為它同樣已經送出過原始碼。", ".sensai/sent.log"),
    ("端點自行掌握", "模型端點由我方指定（Claude Code Router），可接內部或指定供應商；"
     "端點未啟動時靜默停用，不影響開發。", "sensai.endpoint"),
    ("規則不隨工具散布", "規則留在專案 repo 或部門私有 rules repository，"
     "不打包進擴充、不上架 Marketplace。", ".sensai/rules.yaml"),
]
y = Inches(1.8)
for t, b, code in items:
    rect(s, Inches(0.62), y, Inches(12.1), Inches(1.12), fill=WHITE, line_col=LINE)
    rect(s, Inches(0.62), y, Emu(57150), Inches(1.12), fill=TEAL)
    tf = tb(s, Inches(1.0), y + Inches(0.16), Inches(8.6), Inches(0.85))
    para(tf, t, size=14.5, color=NAVY, bold=True, first=True, space_after=4)
    para(tf, b, size=12, color=DARK, space_after=0, line=1.25)
    tf = tb(s, Inches(9.9), y + Inches(0.38), Inches(2.6), Inches(0.4))
    para(tf, code, size=11, color=GREY, first=True, space_after=0, font=MONO)
    y = y + Inches(1.24)
tf = tb(s, Inches(0.62), Inches(6.78), Inches(12.1), Inches(0.3))
para(tf, "※ 未設定任何規則時，sensAI 只檢查語法與型別錯誤，不會做其他推論。",
     size=11.5, color=GREY, first=True, space_after=0)

# ───────────────────────────────── 10 Token 用量
s = new("TOKEN FOOTPRINT", "單次審查的實際用量：為什麼適合地端模型", 10)
tf = tb(s, Inches(0.62), Inches(1.72), Inches(12.1), Inches(0.4))
para(tf, "以下 prompt 大小為實際產生後量測（npm run review -- <檔案> --show-prompt），"
         "非估計值。token 數以程式碼約 3.5〜4 bytes／token 換算。",
     size=13, color=DARK, first=True, space_after=0, line=1.25)

# 左：實測表
rect(s, Inches(0.62), Inches(2.3), Inches(6.55), Inches(3.62), fill=None, line_col=LINE)
tf = tb(s, Inches(0.92), Inches(2.5), Inches(6.0), Inches(0.35))
para(tf, "單次請求的 prompt 大小（實測）", size=13, color=TEAL, bold=True,
     first=True, space_after=0)

# 表頭
tf = tb(s, Inches(1.06), Inches(2.92), Inches(2.9), Inches(0.24))
para(tf, "情境", size=10.5, color=GREY, bold=True, first=True, space_after=0)
tf = tb(s, Inches(4.1), Inches(2.92), Inches(1.3), Inches(0.24), align=PP_ALIGN.RIGHT)
para(tf, "bytes", size=10.5, color=GREY, bold=True, first=True, space_after=0,
     align=PP_ALIGN.RIGHT)
tf = tb(s, Inches(5.55), Inches(2.92), Inches(1.3), Inches(0.24), align=PP_ALIGN.RIGHT)
para(tf, "≈ tokens", size=10.5, color=GREY, bold=True, first=True, space_after=0,
     align=PP_ALIGN.RIGHT)
rect(s, Inches(1.06), Inches(3.2), Inches(5.8), Emu(9525), fill=LINE)

rows = [("C 檔（uart_dma.c ＋ 1 header ＋ 10 條規則）", "20,142", "5.0k〜5.8k", False),
        ("組語檔（uart_dma.s ＋ 內建 ABI 事實）", "12,254", "3.1k〜3.5k", False),
        ("上下文預算滿載（硬上限）", "145,654", "36k〜42k", True)]
y = Inches(3.32)
for i, (name, b, t, hi) in enumerate(rows):
    if hi:
        rect(s, Inches(1.06), y - Inches(0.04), Inches(5.8), Inches(0.52), fill=LIGHT)
    tfx = tb(s, Inches(1.06), y + Inches(0.03), Inches(2.95), Inches(0.44))
    para(tfx, name, size=11.5, color=(NAVY if hi else DARK), bold=hi,
         first=True, space_after=0, line=1.18)
    tfx = tb(s, Inches(4.1), y + Inches(0.05), Inches(1.3), Inches(0.3),
             align=PP_ALIGN.RIGHT)
    para(tfx, b, size=12.5, color=(AMBER if hi else NAVY), bold=True, first=True,
         space_after=0, align=PP_ALIGN.RIGHT)
    tfx = tb(s, Inches(5.55), y + Inches(0.05), Inches(1.3), Inches(0.3),
             align=PP_ALIGN.RIGHT)
    para(tfx, t, size=12.5, color=(AMBER if hi else NAVY), bold=True, first=True,
         space_after=0, align=PP_ALIGN.RIGHT)
    y = y + Inches(0.6)

rect(s, Inches(1.06), Inches(5.18), Inches(5.8), Inches(0.62), fill=NAVY)
tf = tb(s, Inches(1.26), Inches(5.3), Inches(5.4), Inches(0.42))
para(tf, "每次審查送出 1〜2 個請求（兩階段）；連續存檔期間降為 1 個。",
     size=12, color=WHITE, bold=True, first=True, space_after=0, line=1.2)

# 右：地端模型論點
rect(s, Inches(7.45), Inches(2.3), Inches(5.27), Inches(3.62), fill=LIGHT)
tf = tb(s, Inches(7.75), Inches(2.52), Inches(4.67), Inches(3.3))
para(tf, "為什麼這個形狀適合地端模型", size=13, color=TEAL, bold=True,
     first=True, space_after=10)
for i, (t, b) in enumerate([
        ("上界是設定值，不是期望值",
         "contextBudgetBytes 預設 120 KB 是硬上限，超過就截斷。單次請求的最大值"
         "可事先算出 —— 地端 GPU 的 context window 與顯存據此規劃即可。"),
        ("單輪請求，沒有 agentic 迴圈",
         "一次請求換一個回答。沒有 tool use 多輪探索把 token 一層層疊上去，"
         "也就沒有「這次特別貴」的意外。"),
        ("不做全庫檢索",
         "只帶當前檔案與它 #include 到的專案 header。不需要向量資料庫、"
         "不需要全庫索引、不需要 build system 整合。"),
        ("負載隨存檔次數成長，不隨 repo 大小成長",
         "百萬行的專案與一萬行的專案，單次請求大小相同。容量規劃因此可行。")]):
    para(tf, t, size=12, color=NAVY, bold=True, space_after=3,
         space_before=(0 if i == 0 else 9))
    para(tf, b, size=10.5, color=DARK, space_after=0, line=1.22)

# 底部
rect(s, Inches(0.62), Inches(6.08), Inches(12.1), Inches(0.68), fill=NAVY)
tf = tb(s, Inches(0.95), Inches(6.2), Inches(11.5), Inches(0.46))
para(tf, "地端模型的前提是「用量可預測」。sensAI 的用量由設定值決定上界，"
         "不由程式碼庫規模或模型的探索意願決定。",
     size=13, color=WHITE, bold=True, first=True, space_after=0)

# ───────────────────────────────── 11 進度
s = new("STATUS", "目前進度：已可使用，非概念驗證", 11)
stat(s, Inches(0.62), Inches(1.78), Inches(2.9), "14 天", "從構想到可用版本",
     "2026/08/22 → 09/04　45 次提交")
stat(s, Inches(3.72), Inches(1.78), Inches(2.9), "3,603 行", "TypeScript　19 個模組",
     "另有 1,495 行測試程式")
stat(s, Inches(6.82), Inches(1.78), Inches(2.9), "96 / 96", "自動化測試通過",
     "涵蓋差異判定、上下文、過濾、節流")
stat(s, Inches(9.92), Inches(1.78), Inches(2.8), "v0.5.1", "目前版本",
     "已可安裝於 VS Code 使用")

tf = tb(s, Inches(0.62), Inches(3.4), Inches(6.0), Inches(0.35))
para(tf, "已完成", size=15, color=NAVY, bold=True, first=True, space_after=9)
for t in ["存檔自動審查（C 與組語）、手動審查指令",
          "專案 header 自動解析與上下文預算控管",
          "規則引擎、語言過濾、ABI 事實注入",
          "三層節流、審查取消、結果保留不閃爍",
          "釘選與筆記、誤報靜音與匯出報告",
          "隱私排除清單與稽核日誌"]:
    para(tf, "✓　" + t, size=13, color=DARK, space_after=6, line=1.2)

tf = tb(s, Inches(6.9), Inches(3.4), Inches(5.8), Inches(0.35))
para(tf, "已知限制（誠實揭露）", size=15, color=NAVY, bold=True, first=True, space_after=9)
for t in ["多根工作區目前只讀取第一個資料夾",
          "組語巨集不展開；上下文不足時要求模型保守不報",
          "組語審查僅看單檔，抓不到跨檔的 C／組語簽章不一致",
          "只提供意見，不產生也不套用修補程式",
          "尚未累積足夠樣本計算「意見被實際修掉的比例」"]:
    para(tf, "・" + t, size=13, color=GREY, space_after=6, line=1.2)

# ───────────────────────────────── 12 已有數據
s = new("DATA", "目前已有的數據與出處", 12)
tf = tb(s, Inches(0.62), Inches(1.72), Inches(12.1), Inches(0.4))
para(tf, "以下每個數字都可當場複查。現場使用數據（採納率、延遲分布、成本）要等導入後"
         "由稽核日誌自動累積 —— 今天還沒有。",
     size=14, color=DARK, first=True, space_after=0, line=1.25)


def drow(slide, x, y, w, item, value, source, alt=False):
    """三欄資料列：項目／數值／出處。"""
    if alt:
        rect(slide, x, y, w, Inches(0.34), fill=LIGHT)
    tf = tb(slide, x + Inches(0.14), y + Inches(0.07), Inches(1.72), Inches(0.24))
    para(tf, item, size=11.5, color=DARK, first=True, space_after=0)
    tf = tb(slide, x + Inches(1.92), y + Inches(0.06), Inches(2.08), Inches(0.24))
    para(tf, value, size=12, color=NAVY, bold=True, first=True, space_after=0)
    tf = tb(slide, x + Inches(4.06), y + Inches(0.08), Inches(1.72), Inches(0.24))
    para(tf, source, size=10, color=GREY, first=True, space_after=0, font=MONO)


def dhead(slide, x, y, w, text, col=TEAL):
    rect(slide, x, y, Emu(34290), Inches(0.26), fill=col)
    tf = tb(slide, x + Inches(0.14), y + Inches(0.02), w, Inches(0.24))
    para(tf, text, size=11.5, color=col, bold=True, first=True, space_after=0)


LX, RX, CW = Inches(0.62), Inches(6.84), Inches(5.88)

# 左欄：產品規模與規則庫
dhead(s, LX, Inches(2.32), CW, "產品規模與品質")
y = Inches(2.66)
for i, (a, b, c) in enumerate([
        ("程式碼規模", "3,603 行 ／ 19 個模組", "wc -l src/*.ts"),
        ("測試", "96 項，全數通過", "npm test"),
        ("開發期間", "14 天 ／ 45 次提交", "git log"),
        ("目前版本", "v0.5.1", "package.json")]):
    drow(s, LX, y, CW, a, b, c, alt=(i % 2 == 0))
    y = y + Inches(0.34)

dhead(s, LX, y + Inches(0.14), CW, "規則與知識庫")
y = y + Inches(0.48)
for i, (a, b, c) in enumerate([
        ("團隊規則", "10 條（error 8／warning 2）", ".sensai/rules.yaml"),
        ("架構 ABI 事實", "2 組（Andes V5／ARMv7E-M）", "src/abi.ts")]):
    drow(s, LX, y, CW, a, b, c, alt=(i % 2 == 0))
    y = y + Inches(0.34)

dhead(s, LX, y + Inches(0.14), CW, "運作參數（預設值）")
y = y + Inches(0.48)
for i, (a, b, c) in enumerate([
        ("存檔去抖動", "1,000 ms", "sensai.debounceMs"),
        ("上下文上限", "120 KB ／ include 深度 2", "sensai.includeDepth")]):
    drow(s, LX, y, CW, a, b, c, alt=(i % 2 == 0))
    y = y + Inches(0.34)

# 右欄：實際執行觀測
dhead(s, RX, Inches(2.32), CW, "實際執行觀測（樣本數 n＝1）", col=AMBER)
y = Inches(2.66)
for i, (a, b, c) in enumerate([
        ("單檔審查耗時", "65.9 秒", "phal_sys.c 實機"),
        ("附帶 header", "18 個（上下文已截斷）", "同上"),
        ("產生意見", "1 則 error，無誤報", "同上"),
        ("命中缺陷潛伏", "7 個月", "git blame")]):
    drow(s, RX, y, CW, a, b, c, alt=(i % 2 == 0))
    y = y + Inches(0.34)

rect(s, RX, Inches(4.22), CW, Inches(2.0), fill=LIGHT)
tf = tb(s, RX + Inches(0.3), Inches(4.42), CW - Inches(0.6), Inches(1.7))
para(tf, "為什麼右欄只有一筆", size=13, color=NAVY, bold=True, first=True, space_after=7)
para(tf, "稽核日誌（.sensai/sent.log）只在擴充於 VS Code 實際執行時才寫入，"
         "命令列工具不寫。目前僅有一次實機審查的觀測記錄。",
     size=11.5, color=DARK, space_after=8, line=1.25)
para(tf, "導入後每次請求會自動記一筆，含 ts、file、outcome、headers、bytes、"
         "findings、dropped、durationMs —— 屆時右欄可由單點改為分布，"
         "覆蓋率、意見密度、過濾效果與成本也一併算得出來。",
     size=11.5, color=GREY, space_after=0, line=1.25)

tf = tb(s, Inches(0.62), Inches(6.5), Inches(12.1), Inches(0.3))
para(tf, "※ n＝1 的觀測不能拿來推論平均值。本頁左欄是可驗證的計數，右欄是單次觀測 —— "
         "兩者的證據強度不同，報告時分開看。",
     size=11, color=GREY, first=True, space_after=0)

# ───────────────────────────────── 13 部署狀況
s = new("DEPLOYMENT", "部署狀況：已安裝 8 台開發機", 13)
tf = tb(s, Inches(0.62), Inches(1.72), Inches(12.1), Inches(0.4))
para(tf, "sensAI 沒有透過任何統一派送機制推送。這 8 台，是 8 位工程師各自決定安裝的結果。",
     size=14, color=DARK, first=True, space_after=0, line=1.25)

# 左：主數字
rect(s, Inches(0.62), Inches(2.3), Inches(3.62), Inches(3.62), fill=NAVY)
tf = tb(s, Inches(0.92), Inches(2.72), Inches(3.0), Inches(1.5), align=PP_ALIGN.CENTER)
para(tf, "8", size=96, color=WHITE, bold=True, first=True, space_after=0,
     align=PP_ALIGN.CENTER)
rect(s, Inches(1.62), Inches(4.22), Inches(1.6), Emu(22860), fill=TEAL)
tf = tb(s, Inches(0.92), Inches(4.5), Inches(3.0), Inches(1.2), align=PP_ALIGN.CENTER)
para(tf, "台開發機已安裝", size=17, color=WHITE, bold=True, first=True, space_after=8,
     align=PP_ALIGN.CENTER)
para(tf, "來源：VS Code 擴充頁安裝數\n（可於擴充詳細頁當場複查）",
     size=11, color=RGBColor(0xA8, 0xC4, 0xD8), space_after=0, align=PP_ALIGN.CENTER,
     line=1.3)

# 中：這個數字的意義
rect(s, Inches(4.46), Inches(2.3), Inches(4.0), Inches(3.62), fill=None, line_col=LINE)
tf = tb(s, Inches(4.76), Inches(2.52), Inches(3.4), Inches(3.3))
para(tf, "這個數字的意義", size=13, color=TEAL, bold=True, first=True, space_after=10)
for i, (t, b) in enumerate([
        ("全部為主動安裝",
         "未納入任何統一派送。每一台都要工程師自己從 Marketplace 裝上去。"),
        ("零採購成本",
         "自行開發，無授權費，安裝不需要走採購流程。"),
        ("零設定門檻",
         "一個 Initialize Project 指令即建立專案設定，不需要 build system 整合。")]):
    para(tf, t, size=12.5, color=NAVY, bold=True, space_after=3,
         space_before=(0 if i == 0 else 12))
    para(tf, b, size=11, color=GREY, space_after=0, line=1.25)

# 右：部署階梯
rect(s, Inches(8.68), Inches(2.3), Inches(4.04), Inches(3.62), fill=LIGHT)
tf = tb(s, Inches(8.98), Inches(2.52), Inches(3.44), Inches(0.35))
para(tf, "部署進程", size=13, color=TEAL, bold=True, first=True, space_after=0)
stages = [("開發驗證", "1 台", True),
          ("目前", "8 台", True),
          ("試行專案全員", "待展開", False),
          ("部門推廣", "待評估", False)]
y = Inches(2.98)
for name, val, done in stages:
    col = TEAL if done else RGBColor(0xB0, 0xBC, 0xC8)
    rect(s, Inches(8.98), y + Inches(0.12), Inches(0.16), Inches(0.16), fill=col,
         shape=MSO_SHAPE.OVAL)
    tfx = tb(s, Inches(9.32), y + Inches(0.02), Inches(1.9), Inches(0.35))
    para(tfx, name, size=12.5, color=(NAVY if done else GREY), bold=done,
         first=True, space_after=0)
    tfx = tb(s, Inches(11.3), y + Inches(0.03), Inches(1.15), Inches(0.35),
             align=PP_ALIGN.RIGHT)
    para(tfx, val, size=12.5, color=(NAVY if done else GREY), bold=done,
         first=True, space_after=0, align=PP_ALIGN.RIGHT)
    y = y + Inches(0.62)

tf = tb(s, Inches(8.98), Inches(5.56), Inches(3.44), Inches(0.4))
para(tf, "安裝數不等於活躍台數。每週仍在產生審查記錄的台數，"
         "需收回稽核日誌後才算得出。",
     size=10, color=GREY, first=True, space_after=0, line=1.2)

# 底部
rect(s, Inches(0.62), Inches(6.08), Inches(12.1), Inches(0.68), fill=NAVY)
tf = tb(s, Inches(0.95), Inches(6.2), Inches(11.5), Inches(0.46))
para(tf, "下一個里程碑：試行專案全員部署 —— 目標是讓部署台數從「自己找到的人」"
         "變成「整個專案的標準配備」。",
     size=13, color=WHITE, bold=True, first=True, space_after=0)

# ───────────────────────────────── 14 下一步
s = new("NEXT", "下一步與所需支援", 14)
tf = tb(s, Inches(0.62), Inches(1.8), Inches(12.1), Inches(0.35))
para(tf, "建議以一個專案小規模試行，用真實缺陷驗證效益，再決定是否推廣。",
     size=15, color=DARK, first=True, space_after=0)

phases = [("第 1 階段　試行（4 週）",
           ["指定 1〜2 個專案導入，由專案負責人補齊該專案規則",
            "追蹤指標：部署台數，以及每週仍在產生審查記錄的台數",
            "目標：試行專案全員部署，成為專案標準配備"], TEAL),
          ("第 2 階段　擴大（試行後評估）",
           ["建立部門共用規則庫（private rules repository）",
            "把踩過的坑轉成規則，讓同一個錯不會被犯第二次",
            "評估是否納入 CI，於合併前先跑一次"], NAVY2),
          ("所需支援",
           ["核配一組共用模型端點（CCR）與額度，供團隊連線",
            "指定試行專案與該專案的規則維護人",
            "同意規則庫以部門私有 repository 管理"], AMBER)]
x = Inches(0.62)
for t, its, c in phases:
    rect(s, x, Inches(2.5), Inches(3.9), Inches(3.5), fill=WHITE, line_col=LINE)
    rect(s, x, Inches(2.5), Inches(3.9), Inches(0.08), fill=c)
    tf = tb(s, x + Inches(0.28), Inches(2.82), Inches(3.34), Inches(0.4))
    para(tf, t, size=14.5, color=NAVY, bold=True, first=True, space_after=12)
    for it in its:
        para(tf, "・" + it, size=12, color=DARK, space_after=8, line=1.25)
    x = x + Inches(4.1)

rect(s, Inches(0.62), Inches(6.18), Inches(12.1), Inches(0.62), fill=NAVY)
tf = tb(s, Inches(0.95), Inches(6.32), Inches(11.5), Inches(0.4))
para(tf, "一句話：工具已經做好了，真正的價值在規則 —— 需要的是一個試行專案，"
         "把團隊踩過的坑寫下來。",
     size=14, color=WHITE, bold=True, first=True, space_after=0)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
print("saved", OUT, os.path.getsize(OUT), "bytes,", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
